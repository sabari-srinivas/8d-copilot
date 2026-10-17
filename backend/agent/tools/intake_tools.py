"""Complaint text and the messy shift-lead spreadsheet, turned into typed rows."""
import io
import json
import logging
import re
from datetime import date, datetime, timedelta

import pandas as pd
from openpyxl import load_workbook

from agent.models import Complaint, TravelerRow
from agent.tools.llm import generate_structured, load_prompt
from agent.tools.plant_data import get_store
from data.constants import SCENARIO_YEAR

log = logging.getLogger(__name__)

# Short column names the shift leads actually type -> our field names.
COLUMN_ALIASES = {
    "lot#": "lot_id", "lot": "lot_id", "lot no": "lot_id",
    "dt": "date", "date": "date",
    "shift": "shift",
    "opr": "operator_id", "operator": "operator_id",
    "oven peak (c) [from printout]": "reflow_peak_temp_c", "oven peak": "reflow_peak_temp_c",
    "peak": "reflow_peak_temp_c", "peak (c)": "reflow_peak_temp_c",
    "tal sec": "time_above_liquidus_s", "tal": "time_above_liquidus_s",
    "aoi bridge": "aoi_bridges", "aoi bridges": "aoi_bridges", "bridges": "aoi_bridges",
    "remarks": "remarks", "remark": "remarks", "notes": "remarks",
}
NON_NUMERIC = re.compile(r"[^0-9.\-]")


def working_days_between(start: date, end: date) -> int:
    """Signed count of Mon-Fri days strictly after `start` up to and including `end`."""
    sign = 1 if end >= start else -1
    lo, hi = (start, end) if sign == 1 else (end, start)
    days = sum(1 for n in range(1, (hi - lo).days + 1)
               if (lo + timedelta(days=n)).weekday() < 5)
    return sign * days


def _norm_lot(value) -> str:
    text = str(value).strip().upper()
    digits = re.sub(r"\D", "", text)
    return f"L{digits}" if digits else ""


def _norm_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    for fmt in ("%d/%m", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.replace(year=SCENARIO_YEAR).date().isoformat() \
                if fmt == "%d/%m" else parsed.date().isoformat()
        except ValueError:
            continue
    return ""


def _norm_number(value) -> float | None:
    """Returns None for anything the shift lead wrote instead of a number."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = NON_NUMERIC.sub("", str(value))
    try:
        return float(cleaned) if cleaned not in ("", "-", ".") else None
    except ValueError:
        return None


def parse_lot_traveler(xlsx_bytes: bytes) -> list[TravelerRow]:
    try:
        ws = load_workbook(io.BytesIO(xlsx_bytes), data_only=True).active
    except Exception as exc:
        raise ValueError(f"could not read the uploaded spreadsheet: {exc.__class__.__name__}")

    grid = [[c for c in row] for row in ws.iter_rows(values_only=True)]

    # Hunt for the header row: the first row where 2+ cells match a known column name.
    header_idx, mapping = None, {}
    for idx, row in enumerate(grid):
        found = {}
        for col, cell in enumerate(row):
            key = str(cell).strip().lower() if cell is not None else ""
            if key in COLUMN_ALIASES:
                found[COLUMN_ALIASES[key]] = col
        if len(found) >= 2:
            header_idx, mapping = idx, found
            break
    if header_idx is None:
        raise ValueError("could not read the uploaded spreadsheet: no recognisable header row")
    log.info("traveler header found on row %d, %d columns mapped",
             header_idx + 1, len(mapping))

    rows: list[TravelerRow] = []
    for raw in grid[header_idx + 1:]:
        if not any(cell not in (None, "") for cell in raw):
            continue

        def cell(field):
            col = mapping.get(field)
            return raw[col] if col is not None and col < len(raw) else None

        lot_id = _norm_lot(cell("lot_id"))
        if not lot_id:
            continue

        low: list[str] = []
        peak = _norm_number(cell("reflow_peak_temp_c"))
        if peak is None and cell("reflow_peak_temp_c") not in (None, ""):
            low.append("reflow_peak_temp_c")      # e.g. the literal "same as prev"
        tal = _norm_number(cell("time_above_liquidus_s"))
        if tal is None and cell("time_above_liquidus_s") not in (None, ""):
            low.append("time_above_liquidus_s")
        when = _norm_date(cell("date"))
        if not when and cell("date") not in (None, ""):
            low.append("date")

        bridges = _norm_number(cell("aoi_bridges"))
        rows.append(TravelerRow(
            lot_id=lot_id, date=when, shift=str(cell("shift") or "").strip(),
            operator_id=str(cell("operator_id") or "").strip(),
            reflow_peak_temp_c=peak, time_above_liquidus_s=tal,
            aoi_bridges=int(bridges) if bridges is not None else None,
            remarks=str(cell("remarks") or "").strip(), low_confidence_fields=low))
    return rows


def parse_complaint(text: str, complaint_id: str = "") -> Complaint:
    prompt = f"{load_prompt('intake')}\n\nComplaint text:\n{text}"
    parsed = generate_structured(prompt, Complaint, temperature=0.0)
    parsed.text = text
    if complaint_id:
        parsed.complaint_id = complaint_id
    if parsed.lot_id:
        parsed.lot_id = _norm_lot(parsed.lot_id)
    return parsed


def _cell(row: dict, field: str, default: str = "") -> str:
    """A CSV round-trip turns an empty string into NaN (float), not back into "" --
    `pd.read_csv` has no notion of "empty but present" versus "absent" for a text
    column. Every text field read off a complaints.csv row must go through this,
    not a bare `row.get(...)`."""
    value = row.get(field, default)
    return default if value is None or (isinstance(value, float) and pd.isna(value)) else str(value)


def _row_to_complaint(row: dict, today: date) -> Complaint:
    """One `data/out/complaints.csv` row -> a typed `Complaint`. Shared by the inbox
    listing and the single-record lookup so the two can never drift apart on how a
    rejection's fields are parsed."""
    due = _cell(row, "eight_d_due")
    raw_eight_d = _cell(row, "original_eight_d")
    is_rejection = str(row.get("is_rejection", False)).strip().lower() == "true"
    return Complaint(
        complaint_id=row["complaint_id"], customer=row["customer"], product=row["product"],
        lot_id=row["lot_id"], defect=row["defect"], received=_cell(row, "received"),
        eight_d_due=due, status=row.get("status", "open"), text=_cell(row, "text"),
        due_in_working_days=working_days_between(today, date.fromisoformat(due))
        if due else None,
        is_rejection=is_rejection, objection=_cell(row, "objection"),
        original_eight_d=json.loads(raw_eight_d) if raw_eight_d else None)


def list_complaints(today: date | None = None) -> list[Complaint]:
    """The inbox. Parsed fields come from the dataset, so no model call per page load
    (PRD 7 screen 0: 'cached after first run')."""
    today = today or date.today()
    return [_row_to_complaint(row, today) for row in get_store().complaints().to_dict("records")]


def get_complaint_record(complaint_id: str, today: date | None = None) -> Complaint | None:
    """Deterministic, CSV-backed single-complaint lookup -- used by
    agent/orchestrator.py to learn whether the complaint behind a live investigation is
    a rejection, and if so its objection and the 8D it is rejecting. This deliberately
    does NOT go through `parse_complaint` (the LLM path, used for `customer`/`product`/
    `lot_id` parsed from free-form text): whether a complaint is a rejection, and the
    full D1-D8 text it is rejecting, are facts about the complaint record, not
    something a model should reconstruct from a snippet of text on every run."""
    today = today or date.today()
    rows = get_store().complaints().to_dict("records")
    row = next((r for r in rows if r["complaint_id"] == complaint_id), None)
    return _row_to_complaint(row, today) if row is not None else None
