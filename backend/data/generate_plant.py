"""
generate_plant.py — synthetic SMT plant data for the 8D Copilot demo.

ALL DATA IS SYNTHETIC. Value ranges are shaped on public SMT references
(SAC305 lead-free profile: preheat ramp 1-3 C/s, soak 150-180 C for 60-120 s,
peak 235-250 C, ~217 C melting point). The true root cause is planted on purpose.

Planted truth  : Line 2, after WO-388 (reflow oven zone 7 heater replaced on
                 2026-10-08, profile re-tuned, verification NOT done), peak temp
                 drifts ~+8 C (245 -> 253, above the 250 C limit) and time above
                 liquidus rises. Solder bridge rate ~6x. Lots L2609-L2614.
Decoy 1 (H1)   : stencil wear. Line 2 got a NEW stencil (WO-386, 2026-10-07),
                 so the bad lots have LOW print counts -> rejected.
Decoy 2 (H3)   : room humidity. Random 45-60 %, unrelated -> rejected.
Decoy 3 (H4)   : operator OP-17 joined Line 2 on 2026-10-06, close to WO-388,
                 so the data can't separate them cleanly -> inconclusive.

Run: python -m data.generate_plant   (writes data/out/)
"""
import argparse
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from data.constants import (DRIFT_LOT_IDS, OP17_L1_SHARE, OP17_START, PRODUCT, SEED,
                            START_DATE, WO386_DATE, WO388_DATE)

ap = argparse.ArgumentParser()
ap.add_argument("--seed", type=int, default=SEED)
ap.add_argument("--op17-l1-share", type=float, default=OP17_L1_SHARE)
args = ap.parse_args()
rng = np.random.default_rng(args.seed)

OUT = Path(__file__).parent / "out"
OUT.mkdir(parents=True, exist_ok=True)

START = pd.Timestamp(START_DATE)
WO388_DATE = pd.Timestamp(WO388_DATE)
WO386_DATE = pd.Timestamp(WO386_DATE)
OP17_START = pd.Timestamp(OP17_START)

# ---------- lots ----------
rows = []
lot_ids = list(range(2551, 2615))              # 64 lots
dates = pd.date_range(START, "2026-10-13", periods=len(lot_ids)).normalize()
for i, (lot, d) in enumerate(zip(lot_ids, dates)):
    lot_id = f"L{lot}"
    after_wo388 = d > pd.Timestamp(WO388_DATE)
    if lot_id in DRIFT_LOT_IDS:
        line = "L2"                      # Line 1 ran ECU-PCB-9 on these days
    elif after_wo388:
        line = "L1"                      # keep the post-WO-388 L2 window exactly the drift lots
    else:
        line = "L1" if i % 2 == 0 else "L2"
    shift = "ABC"[i % 3]
    # operators
    if line == "L1":
        op = rng.choice(["OP-04", "OP-09", "OP-17"],
                        p=[0.45, 0.55 - args.op17_l1_share, args.op17_l1_share])
    elif d >= OP17_START:
        op = "OP-17" if shift in "AB" else "OP-22"
    else:
        op = rng.choice(["OP-11", "OP-22"], p=[0.6, 0.4])
    # stencil
    if line == "L1":
        stencil = "ST-03A"
    else:
        stencil = "ST-07B" if d >= pd.Timestamp("2026-10-07") else "ST-05A"
    paste_lot = f"PL-{2601 + (d - START).days // 7:04d}"
    qty = int(rng.integers(480, 521))
    rows.append(dict(lot_id=lot_id, line=line, shift=shift, date=d.date(),
                     product=PRODUCT, qty=qty, paste_lot=paste_lot,
                     stencil_id=stencil, operator_id=op))
lots = pd.DataFrame(rows)

post = lots[(lots.line == "L2") & (pd.to_datetime(lots.date) > pd.Timestamp(WO388_DATE))]
assert sorted(post.lot_id) == sorted(DRIFT_LOT_IDS), (
    f"post-WO-388 Line 2 lots {sorted(post.lot_id)} != drift lots {sorted(DRIFT_LOT_IDS)}")

# ---------- process ----------
proc = []
stencil_counts = {"ST-03A": 18000, "ST-05A": 41000, "ST-07B": 0}
for r in lots.itertuples():
    drift = r.line == "L2" and pd.Timestamp(r.date) > WO388_DATE
    stencil_counts[r.stencil_id] += int(r.qty / 2) + int(rng.integers(0, 40))  # 2-up panels
    proc.append(dict(
        lot_id=r.lot_id,
        reflow_peak_temp_c=round(rng.normal(253.2 if drift else 245.0, 1.4), 1),
        time_above_liquidus_s=int(rng.normal(82 if drift else 62, 4)),
        preheat_ramp_c_per_s=round(rng.normal(1.6, 0.15), 2),
        soak_time_s=int(rng.normal(85, 6)),
        conveyor_speed_cm_min=round(rng.normal(78 if r.line == "L1" else 80, 1.0), 1),
        paste_age_h=round(float(rng.uniform(2, 20)), 1),
        room_humidity_pct=round(float(rng.uniform(45, 60)), 1),
        stencil_print_count=stencil_counts[r.stencil_id],
        squeegee_pressure_kg=round(rng.normal(6.0, 0.2), 2),
    ))
process = pd.DataFrame(proc)

# ---------- defects (AOI + customer returns) ----------
def bridges(r):
    if r.lot_id == "L2611":
        return 14                              # matches the customer complaint
    drift = r.line == "L2" and pd.Timestamp(r.date) > WO388_DATE
    return int(rng.poisson(r.qty * (0.026 if drift else 0.004)))
lots["defects_solder_bridge"] = [bridges(r) for r in lots.itertuples()]
lots["defects_other"] = rng.poisson(1.5, len(lots))
lots["defect_rate"] = ((lots.defects_solder_bridge + lots.defects_other) / lots.qty).round(4)
lots["ship_status"] = np.where(lots.lot_id.isin(["L2613", "L2614"]), "in_warehouse", "shipped")

# ---------- maintenance ----------
maint = pd.DataFrame([
    ("WO-371", "L1", "2026-09-16", "Pick-and-place head 2", "Nozzle set replaced (wear)", "closed", "profile_check_na"),
    ("WO-374", "L2", "2026-09-19", "Stencil printer", "Squeegee blade replaced", "closed", "n/a"),
    ("WO-377", "L1", "2026-09-23", "Reflow oven", "Quarterly PM: conveyor chain lubrication, fan check", "closed", "profile_verified"),
    ("WO-379", "L2", "2026-09-25", "AOI station", "Camera recalibration", "closed", "n/a"),
    ("WO-381", "L1", "2026-09-29", "Nitrogen generator", "Filter change", "closed", "n/a"),
    ("WO-383", "L2", "2026-10-02", "Reflow oven", "Quarterly PM: conveyor chain lubrication, fan check", "closed", "profile_verified"),
    ("WO-386", "L2", "2026-10-07", "Stencil printer", "New stencil ST-07B installed for ECU-PCB-7 (replaces ST-05A)", "closed", "first_article_ok"),
    ("WO-388", "L2", "2026-10-08", "Reflow oven zone 7", "Zone 7 heater element replaced, profile re-tuned by vendor", "closed", "profile_verification_pending"),
    ("WO-390", "L1", "2026-10-10", "Pick-and-place feeder bank", "Feeder 14 jam cleared", "closed", "n/a"),
    ("WO-392", "L2", "2026-10-12", "Conveyor", "Board stop sensor cleaned", "closed", "n/a"),
], columns=["wo_id", "line", "date", "asset", "description", "status", "post_check"])

# ---------- complaints (inbox) ----------
complaints = pd.DataFrame([
    ("C-2026-031", "2026-10-15", "Hoshino Denso (fictional)", "ECU-PCB-7", "L2611", "solder_bridge",
     "Solder bridging found on ECU-PCB-7, lot 2611, 14 of 500 units at our incoming/ICT. "
     "Please contain immediately and send 8D within 10 working days.", "2026-10-29", "open"),
    ("C-2026-032", "2026-10-14", "Kaveri Motors (fictional)", "ECU-PCB-9", "L9142", "tombstoning",
     "3 boards with tombstoned 0402 capacitors at C112, lot 9142. Request 8D.", "2026-10-28", "open"),
    ("C-2026-024", "2026-07-02", "Hoshino Denso (fictional)", "ECU-PCB-7", "L2389", "tombstoning",
     "Tombstoning on 0402 resistors, lot 2389.", "2026-07-16", "closed"),
], columns=["complaint_id", "received", "customer", "product", "lot_id", "defect",
            "text", "eight_d_due", "status"])

lots.to_csv(OUT / "lots.csv", index=False)
process.to_csv(OUT / "process.csv", index=False)
maint.to_csv(OUT / "maintenance.csv", index=False)
complaints.to_csv(OUT / "complaints.csv", index=False)

# ---------- messy lot traveler (the "no MES" input) ----------
wb = Workbook()
ws = wb.active
ws.title = "Line2 Oct wk41"
ws["A1"] = "SREE ELECTRONICS (fictional) - SMT LINE 2 LOT TRAVELER - ECU-PCB-7"
ws["A1"].font = Font(bold=True)
ws["A2"] = "Synthetic demo file. Filled by shift leads, not validated."
ws.append([])
ws.append(["Lot#", "Dt", "Shift", "Opr", "Paste lot", "Stencil", "Oven peak (C) [from printout]",
           "TAL sec", "AOI bridge", "Remarks"])
for c in ws[4]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBDD")
remarks = {
    "L2609": "oven done after maint, ok to run per vendor",
    "L2610": "AOI few bridges U7, reworked",
    "L2611": "bridges near U7 QFP fine pitch - reworked 9, rest shipped",
    "L2612": "",
    "L2613": "profile printout missing, used last one",
    "L2614": "bridging again. informed QE?",
}
m = lots.merge(process, on="lot_id")
for r in m[m.line == "L2"].tail(9).itertuples():
    peak = "same as prev" if r.lot_id == "L2613" else r.reflow_peak_temp_c
    ws.append([r.lot_id.replace("L", ""), pd.Timestamp(r.date).strftime("%d/%m"), r.shift,
               r.operator_id, r.paste_lot, r.stencil_id, peak, r.time_above_liquidus_s,
               r.defects_solder_bridge, remarks.get(r.lot_id, "")])
for col, w in zip("ABCDEFGHIJ", [8, 7, 6, 7, 10, 9, 26, 8, 10, 48]):
    ws.column_dimensions[col].width = w
wb.save(OUT / "lot_traveler_line2_wk41.xlsx")

print(f"wrote {len(lots)} lots, {len(maint)} work orders, {len(complaints)} complaints to {OUT}")
